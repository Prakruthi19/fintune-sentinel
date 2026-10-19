"""Turn an uploaded statement (PDF, Excel, Word, CSV or text) into plain lines for parse_statement.

Files are read in memory and never written to disk, so "nothing saved on our servers" holds.
"""
import csv
import io
import os

MAX_BYTES = 10 * 1024 * 1024
EXTENSIONS = (".pdf", ".xlsx", ".xlsm", ".docx", ".csv", ".tsv", ".txt", ".md")


class DocumentError(ValueError):
    """A message safe to show the user as-is."""


def _cell(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return "" if value is None else str(value)


def _pdf(data):
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            reader.decrypt("")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as err:
        raise DocumentError("Couldn't open this PDF. If it's password-protected, remove the password first.") from err


def _excel(data):
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as err:
        raise DocumentError("Couldn't open this Excel file. Save it as .xlsx and try again.") from err
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [_cell(v) for v in row if v not in (None, "")]
            if cells:
                lines.append("\t".join(cells))
    return "\n".join(lines)


def _docx(data):
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as err:
        raise DocumentError("Couldn't open this Word file. Save it as .docx and try again.") from err
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append("\t".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


def _text(data, delimiter=None):
    text = data.decode("utf-8-sig", errors="replace")
    if delimiter:
        text = "\n".join("\t".join(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter))
    return text


def extract_text(filename, data):
    ext = os.path.splitext(filename or "")[1].lower()
    if ext not in EXTENSIONS:
        raise DocumentError("Upload a PDF, Excel (.xlsx), Word (.docx), CSV or text file.")
    if len(data) > MAX_BYTES:
        raise DocumentError("That file is over 10 MB. Upload just the financial statements.")
    if ext == ".pdf":
        text = _pdf(data)
        if not text.strip():
            raise DocumentError("This PDF has no readable text; it looks like a scan. "
                                "Export it from your accounting software, or type the numbers in.")
        return text
    if ext in (".xlsx", ".xlsm"):
        return _excel(data)
    if ext == ".docx":
        return _docx(data)
    return _text(data, {".csv": ",", ".tsv": "\t"}.get(ext))
