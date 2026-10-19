import io

import pytest
from fastapi.testclient import TestClient

import app
from sentinel.documents import DocumentError, extract_text
from tests.test_app import FakeLlm

ROWS = [("Line item", 2024, 2023), ("Revenue", 510000, 450000), ("Cost of goods sold", 204000, 180000),
        ("Net income", 48000, 39000), ("Total current assets", 140000, 120000)]
EXPECTED = {"revenue": 510000, "revenue_prior": 450000, "cogs": 204000, "net_income": 48000, "current_assets": 140000}


def xlsx():
    from openpyxl import Workbook

    wb = Workbook()
    for row in ROWS:
        wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def docx():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Coffee shop, income statement")
    table = doc.add_table(rows=0, cols=3)
    for row in ROWS:
        cells = table.add_row().cells
        for c, v in zip(cells, row):
            c.text = f"{v:,}" if isinstance(v, int) else v
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def csv_bytes():
    return "\n".join(",".join(f'"{v:,}"' if isinstance(v, int) else v for v in row) for row in ROWS).encode()


def test_formats_parse_to_the_same_fields():
    from sentinel.health import parse_statement

    for name, data in (("s.xlsx", xlsx()), ("s.docx", docx()), ("s.csv", csv_bytes())):
        assert parse_statement(extract_text(name, data)) == EXPECTED, name


def test_rejects_unsupported_and_scanned():
    with pytest.raises(DocumentError, match="Upload a PDF"):
        extract_text("photo.jpg", b"x")
    with pytest.raises(DocumentError, match="Couldn't open this PDF"):
        extract_text("broken.pdf", b"not a pdf")


def test_upload_endpoint():
    client = TestClient(app.create_app(FakeLlm([]), "fake.gguf"))
    r = client.post("/api/upload", files={"file": ("pl.xlsx", xlsx())})
    assert r.status_code == 200 and r.json()["fields"] == EXPECTED
    r = client.post("/api/upload", files={"file": ("pl.exe", b"x")})
    assert r.status_code == 422 and "Upload a PDF" in r.json()["detail"]
